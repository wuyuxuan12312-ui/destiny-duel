# -*- coding: utf-8 -*-
"""
Generate CardGame_Cards.xlsx with 4 sheets:
1. Cards_Base (12 cards)
2. Cards_PVP (48 cards)
3. Cards_PVE (48 cards)
4. Card_Effects (Dictionary of the 7 effect types)
"""
import os
import sys
import openpyxl

if hasattr(sys.stdout, 'reconfigure'):
    try:
        sys.stdout.reconfigure(encoding='utf-8', errors='replace')
        sys.stderr.reconfigure(encoding='utf-8', errors='replace')
    except Exception:
        pass
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

PROJECT_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUTPUT_PATH = os.path.join(PROJECT_ROOT, 'config', 'CardGame_Cards.xlsx')

wb = openpyxl.Workbook()
# remove default sheet
default_sheet = wb.active
wb.remove(default_sheet)

HEADER_FILL = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
HEADER_FONT = Font(name="Microsoft YaHei", size=11, bold=True, color="FFFFFF")
CELL_FONT = Font(name="Microsoft YaHei", size=10)
ALIGN_CENTER = Alignment(horizontal="center", vertical="center")
ALIGN_LEFT = Alignment(horizontal="left", vertical="center")
BORDER_THIN = Border(
    left=Side(style='thin', color='E2E8F0'),
    right=Side(style='thin', color='E2E8F0'),
    top=Side(style='thin', color='E2E8F0'),
    bottom=Side(style='thin', color='E2E8F0')
)

COLUMNS = [
    ("id", "卡牌ID", 20),
    ("name", "卡牌名称", 16),
    ("description", "效果描述", 36),
    ("cost", "消耗能量", 10),
    ("rarity", "品质", 10),
    ("type", "类型", 12),
    ("pool_type", "卡池类型", 12),
    ("upgradeable", "是否可升级", 12),
    ("level", "初始等级", 10),
    ("effect_type", "主要效果", 12),
    ("damage", "伤害数值", 10),
    ("heal", "治疗数值", 10),
    ("shield", "护盾数值", 10),
    ("draw", "抽牌数量", 10),
    ("energy", "能量变化", 10),
    ("status", "状态效果", 18),
    ("buff", "增益效果", 22),
    ("target", "目标类型", 10),
    ("hit_count", "攻击段数", 10),
    ("self_damage", "反噬伤害", 10),
    ("proc_chance", "触发概率", 10),
    ("condition", "触发条件", 22),
    ("condition_param", "条件参数", 10),
    ("condition_bonus", "条件追加数值", 12),
    ("tags", "流派标签", 16)
]

# 1. Base Cards (12)
BASE_CARDS = [
    {
        "id": "base_quick_attack", "name": "快速攻击", "description": "消耗1点能量，造成 8 点快速物理伤害。",
        "cost": 1, "rarity": "common", "type": "attack", "pool_type": "Base", "upgradeable": 0, "level": 1,
        "effect_type": "damage", "damage": 8, "heal": 0, "shield": 0, "draw": 0, "energy": 0,
        "status": "", "buff": "", "target": "enemy", "hit_count": 1, "self_damage": 0, "proc_chance": 0,
        "condition": "", "condition_param": 0, "condition_bonus": 0, "tags": "attack"
    },
    {
        "id": "base_heavy_strike", "name": "重击", "description": "消耗3点能量，蓄力发动猛击，造成 22 点重型伤害。",
        "cost": 3, "rarity": "rare", "type": "attack", "pool_type": "Base", "upgradeable": 0, "level": 1,
        "effect_type": "damage", "damage": 22, "heal": 0, "shield": 0, "draw": 0, "energy": 0,
        "status": "", "buff": "", "target": "enemy", "hit_count": 1, "self_damage": 0, "proc_chance": 0,
        "condition": "", "condition_param": 0, "condition_bonus": 0, "tags": "attack"
    },
    {
        "id": "base_small_shield", "name": "小型护盾", "description": "消耗1点能量，架起坚实屏障，自身获得 15 点护盾。",
        "cost": 1, "rarity": "common", "type": "defense", "pool_type": "Base", "upgradeable": 0, "level": 1,
        "effect_type": "shield", "damage": 0, "heal": 0, "shield": 15, "draw": 0, "energy": 0,
        "status": "", "buff": "", "target": "self", "hit_count": 1, "self_damage": 0, "proc_chance": 0,
        "condition": "", "condition_param": 0, "condition_bonus": 0, "tags": "defense"
    },
    {
        "id": "base_heal", "name": "治疗术", "description": "消耗2点能量，凝聚温和圣光，为自身恢复 15 点生命值。",
        "cost": 2, "rarity": "common", "type": "heal", "pool_type": "Base", "upgradeable": 0, "level": 1,
        "effect_type": "heal", "damage": 0, "heal": 15, "shield": 0, "draw": 0, "energy": 0,
        "status": "", "buff": "", "target": "self", "hit_count": 1, "self_damage": 0, "proc_chance": 0,
        "condition": "", "condition_param": 0, "condition_bonus": 0, "tags": "heal"
    },
    {
        "id": "base_flame_flask", "name": "火焰瓶", "description": "消耗2点能量，投掷爆燃火瓶，造成 12 点伤害并附加 2 层灼烧。",
        "cost": 2, "rarity": "rare", "type": "attack", "pool_type": "Base", "upgradeable": 0, "level": 1,
        "effect_type": "damage", "damage": 12, "heal": 0, "shield": 0, "draw": 0, "energy": 0,
        "status": "burn:2", "buff": "", "target": "enemy", "hit_count": 1, "self_damage": 0, "proc_chance": 0,
        "condition": "", "condition_param": 0, "condition_bonus": 0, "tags": "fire,attack"
    },
    {
        "id": "base_ice_crystal", "name": "冰晶", "description": "消耗2点能量，射出极寒棱晶，造成 10 点伤害，50%几率冻结目标1回合。",
        "cost": 2, "rarity": "rare", "type": "attack", "pool_type": "Base", "upgradeable": 0, "level": 1,
        "effect_type": "damage", "damage": 10, "heal": 0, "shield": 0, "draw": 0, "energy": 0,
        "status": "freeze:1", "buff": "", "target": "enemy", "hit_count": 1, "self_damage": 0, "proc_chance": 0.5,
        "condition": "chance", "condition_param": 0.5, "condition_bonus": 0, "tags": "ice,attack"
    },
    {
        "id": "base_poison_blade", "name": "毒刃", "description": "消耗2点能量，淬毒刀刃突刺，造成 8 点伤害并附加 3 层剧毒。",
        "cost": 2, "rarity": "rare", "type": "attack", "pool_type": "Base", "upgradeable": 0, "level": 1,
        "effect_type": "damage", "damage": 8, "heal": 0, "shield": 0, "draw": 0, "energy": 0,
        "status": "poison:3", "buff": "", "target": "enemy", "hit_count": 1, "self_damage": 0, "proc_chance": 0,
        "condition": "", "condition_param": 0, "condition_bonus": 0, "tags": "poison,attack"
    },
    {
        "id": "base_energy_surge", "name": "能量爆发", "description": "消耗0点能量，引燃体内魔能核心，立即获得 3 点能量。",
        "cost": 0, "rarity": "epic", "type": "resource", "pool_type": "Base", "upgradeable": 0, "level": 1,
        "effect_type": "energy", "damage": 0, "heal": 0, "shield": 0, "draw": 0, "energy": 3,
        "status": "", "buff": "", "target": "self", "hit_count": 1, "self_damage": 0, "proc_chance": 0,
        "condition": "", "condition_param": 0, "condition_bonus": 0, "tags": "defense"
    },
    {
        "id": "base_focus", "name": "专注", "description": "消耗1点能量，屏息凝神，抽取 2 张手牌，并使下一次攻击伤害 +5 点。",
        "cost": 1, "rarity": "rare", "type": "resource", "pool_type": "Base", "upgradeable": 0, "level": 1,
        "effect_type": "draw", "damage": 0, "heal": 0, "shield": 0, "draw": 2, "energy": 0,
        "status": "", "buff": "damage+5", "target": "self", "hit_count": 1, "self_damage": 0, "proc_chance": 0,
        "condition": "", "condition_param": 0, "condition_bonus": 0, "tags": "attack"
    },
    {
        "id": "base_defensive_stance", "name": "防御姿态", "description": "消耗2点能量，沉下重心摆出防守步法，本回合受到的伤害降低 40%。",
        "cost": 2, "rarity": "rare", "type": "defense", "pool_type": "Base", "upgradeable": 0, "level": 1,
        "effect_type": "buff", "damage": 0, "heal": 0, "shield": 0, "draw": 0, "energy": 0,
        "status": "", "buff": "damage_reduce:0.4", "target": "self", "hit_count": 1, "self_damage": 0, "proc_chance": 0,
        "condition": "", "condition_param": 0, "condition_bonus": 0, "tags": "defense"
    },
    {
        "id": "base_dual_slash", "name": "双刃斩", "description": "消耗2点能量，双刀疯狂交叉斩击造成 18 点伤害，自身承受 3 点反噬伤害。",
        "cost": 2, "rarity": "rare", "type": "attack", "pool_type": "Base", "upgradeable": 0, "level": 1,
        "effect_type": "damage", "damage": 18, "heal": 0, "shield": 0, "draw": 0, "energy": 0,
        "status": "", "buff": "", "target": "enemy", "hit_count": 1, "self_damage": 3, "proc_chance": 0,
        "condition": "", "condition_param": 0, "condition_bonus": 0, "tags": "attack"
    },
    {
        "id": "base_purify", "name": "净化", "description": "消耗1点能量，释放清澈灵光，解除自身持有的所有负面异常状态（Debuff）。",
        "cost": 1, "rarity": "common", "type": "skill", "pool_type": "Base", "upgradeable": 0, "level": 1,
        "effect_type": "status", "damage": 0, "heal": 0, "shield": 0, "draw": 0, "energy": 0,
        "status": "remove_negative", "buff": "", "target": "self", "hit_count": 1, "self_damage": 0, "proc_chance": 0,
        "condition": "", "condition_param": 0, "condition_bonus": 0, "tags": "heal"
    }
]

# 2. PVP Cards (48: 12 Attack, 12 Defense, 12 Resource, 12 Control)
PVP_CARDS = [
    # Attack (12)
    {
        "id": "pvp_probe_strike", "name": "试探攻击", "description": "造成 8 点伤害。若目标本回合获得过护盾，则追加 8 点破绽伤害。",
        "cost": 1, "rarity": "common", "type": "attack", "pool_type": "PVP", "upgradeable": 0, "level": 1,
        "effect_type": "damage", "damage": 8, "heal": 0, "shield": 0, "draw": 0, "energy": 0,
        "status": "", "buff": "", "target": "enemy", "hit_count": 1, "self_damage": 0, "proc_chance": 0,
        "condition": "target_shield_gained_this_turn", "condition_param": 0, "condition_bonus": 8, "tags": "attack"
    },
    {
        "id": "pvp_shield_breaker", "name": "破盾刺击", "description": "造成 14 点精准伤害，50% 伤害无视敌方护盾直接穿透扣除生命。",
        "cost": 2, "rarity": "rare", "type": "attack", "pool_type": "PVP", "upgradeable": 0, "level": 1,
        "effect_type": "damage", "damage": 14, "heal": 0, "shield": 0, "draw": 0, "energy": 0,
        "status": "", "buff": "shield_penetration:0.5", "target": "enemy", "hit_count": 1, "self_damage": 0, "proc_chance": 0,
        "condition": "", "condition_param": 0, "condition_bonus": 0, "tags": "attack"
    },
    {
        "id": "pvp_flame_burst", "name": "火焰冲击", "description": "凝聚爆烈热浪，造成 22 点烈焰伤害，并附加 2 层灼烧状态。",
        "cost": 3, "rarity": "rare", "type": "attack", "pool_type": "PVP", "upgradeable": 0, "level": 1,
        "effect_type": "damage", "damage": 22, "heal": 0, "shield": 0, "draw": 0, "energy": 0,
        "status": "burn:2", "buff": "", "target": "enemy", "hit_count": 1, "self_damage": 0, "proc_chance": 0,
        "condition": "", "condition_param": 0, "condition_bonus": 0, "tags": "fire,attack"
    },
    {
        "id": "pvp_thunder_strike", "name": "雷霆一击", "description": "引天雷轰击目标造成 18 点伤害，有 30% 概率触发电涌额外追加 10 点伤害。",
        "cost": 3, "rarity": "rare", "type": "attack", "pool_type": "PVP", "upgradeable": 0, "level": 1,
        "effect_type": "damage", "damage": 18, "heal": 0, "shield": 0, "draw": 0, "energy": 0,
        "status": "", "buff": "", "target": "enemy", "hit_count": 1, "self_damage": 0, "proc_chance": 0.3,
        "condition": "chance", "condition_param": 0.3, "condition_bonus": 10, "tags": "attack"
    },
    {
        "id": "pvp_double_strike", "name": "双重打击", "description": "极速发动连续两次打击，每次造成 10 点伤害（共计 20 点）。",
        "cost": 3, "rarity": "rare", "type": "attack", "pool_type": "PVP", "upgradeable": 0, "level": 1,
        "effect_type": "damage", "damage": 10, "heal": 0, "shield": 0, "draw": 0, "energy": 0,
        "status": "", "buff": "", "target": "enemy", "hit_count": 2, "self_damage": 0, "proc_chance": 0,
        "condition": "", "condition_param": 0, "condition_bonus": 0, "tags": "attack"
    },
    {
        "id": "pvp_life_burn", "name": "生命燃烧", "description": "以自身精血献祭，造成 18 点狂暴伤害，自身受到 5 点反噬伤害。",
        "cost": 2, "rarity": "rare", "type": "attack", "pool_type": "PVP", "upgradeable": 0, "level": 1,
        "effect_type": "damage", "damage": 18, "heal": 0, "shield": 0, "draw": 0, "energy": 0,
        "status": "", "buff": "", "target": "enemy", "hit_count": 1, "self_damage": 5, "proc_chance": 0,
        "condition": "", "condition_param": 0, "condition_bonus": 0, "tags": "attack"
    },
    {
        "id": "pvp_predator", "name": "追猎", "description": "造成 15 点狩猎伤害。若目标本回合内进行过生命治疗，则额外追加 10 点破愈伤害。",
        "cost": 3, "rarity": "rare", "type": "attack", "pool_type": "PVP", "upgradeable": 0, "level": 1,
        "effect_type": "damage", "damage": 15, "heal": 0, "shield": 0, "draw": 0, "energy": 0,
        "status": "", "buff": "", "target": "enemy", "hit_count": 1, "self_damage": 0, "proc_chance": 0,
        "condition": "target_healed_this_turn", "condition_param": 0, "condition_bonus": 10, "tags": "attack"
    },
    {
        "id": "pvp_heartseeker", "name": "穿心箭", "description": "射出破坚锐矢造成 12 点伤害，并击碎削减目标当前拥有的 50% 护盾值。",
        "cost": 2, "rarity": "rare", "type": "attack", "pool_type": "PVP", "upgradeable": 0, "level": 1,
        "effect_type": "damage", "damage": 12, "heal": 0, "shield": 0, "draw": 0, "energy": 0,
        "status": "", "buff": "reduce_target_shield:0.5", "target": "enemy", "hit_count": 1, "self_damage": 0, "proc_chance": 0,
        "condition": "", "condition_param": 0, "condition_bonus": 0, "tags": "attack"
    },
    {
        "id": "pvp_shadow_strike", "name": "影袭", "description": "隐匿暗影中斩出 12 点伤害，并随机迫使敌方弃置 1 张手牌。",
        "cost": 2, "rarity": "epic", "type": "attack", "pool_type": "PVP", "upgradeable": 0, "level": 1,
        "effect_type": "damage", "damage": 12, "heal": 0, "shield": 0, "draw": 0, "energy": 0,
        "status": "", "buff": "discard_enemy_random:1", "target": "enemy", "hit_count": 1, "self_damage": 0, "proc_chance": 0,
        "condition": "", "condition_param": 0, "condition_bonus": 0, "tags": "attack"
    },
    {
        "id": "pvp_berserk_stance", "name": "狂战姿态", "description": "进入嗜血狂化：造成的伤害增加 8 点，但自身本回合受到的所有伤害额外增加 5 点。",
        "cost": 2, "rarity": "rare", "type": "attack", "pool_type": "PVP", "upgradeable": 0, "level": 1,
        "effect_type": "buff", "damage": 0, "heal": 0, "shield": 0, "draw": 0, "energy": 0,
        "status": "", "buff": "damage+8,incoming_damage+5", "target": "self", "hit_count": 1, "self_damage": 0, "proc_chance": 0,
        "condition": "", "condition_param": 0, "condition_bonus": 0, "tags": "attack"
    },
    {
        "id": "pvp_execution", "name": "处决", "description": "挥动巨刃造成 30 点斩击伤害。若目标当前生命值低于 40%，则额外追加 10 点终结伤害。",
        "cost": 5, "rarity": "epic", "type": "attack", "pool_type": "PVP", "upgradeable": 0, "level": 1,
        "effect_type": "damage", "damage": 30, "heal": 0, "shield": 0, "draw": 0, "energy": 0,
        "status": "", "buff": "", "target": "enemy", "hit_count": 1, "self_damage": 0, "proc_chance": 0,
        "condition": "target_hp_below_40", "condition_param": 0.4, "condition_bonus": 10, "tags": "attack"
    },
    {
        "id": "pvp_chain_slash", "name": "连锁斩", "description": "瞬身二段连环斩击，先击造成 10 点伤害，紧接着追加 20 点重击（总计造成 30 点伤害）。",
        "cost": 4, "rarity": "epic", "type": "attack", "pool_type": "PVP", "upgradeable": 0, "level": 1,
        "effect_type": "damage", "damage": 10, "heal": 0, "shield": 0, "draw": 0, "energy": 0,
        "status": "", "buff": "", "target": "enemy", "hit_count": 1, "self_damage": 0, "proc_chance": 0,
        "condition": "always", "condition_param": 0, "condition_bonus": 20, "tags": "attack"
    },

    # Defense (12)
    {
        "id": "pvp_temp_rampart", "name": "临时壁垒", "description": "应急架设城墙石垒，自身立即获得 25 点护盾。",
        "cost": 2, "rarity": "common", "type": "defense", "pool_type": "PVP", "upgradeable": 0, "level": 1,
        "effect_type": "shield", "damage": 0, "heal": 0, "shield": 25, "draw": 0, "energy": 0,
        "status": "", "buff": "", "target": "self", "hit_count": 1, "self_damage": 0, "proc_chance": 0,
        "condition": "", "condition_param": 0, "condition_bonus": 0, "tags": "defense"
    },
    {
        "id": "pvp_iron_armor", "name": "钢铁护甲", "description": "重铸合金战甲覆盖全身，自身立即获得 40 点厚实护盾。",
        "cost": 3, "rarity": "rare", "type": "defense", "pool_type": "PVP", "upgradeable": 0, "level": 1,
        "effect_type": "shield", "damage": 0, "heal": 0, "shield": 40, "draw": 0, "energy": 0,
        "status": "", "buff": "", "target": "self", "hit_count": 1, "self_damage": 0, "proc_chance": 0,
        "condition": "", "condition_param": 0, "condition_bonus": 0, "tags": "defense"
    },
    {
        "id": "pvp_counter_shield", "name": "反击盾", "description": "架起反伤荆棘盾获得 20 点护盾，本回合受到敌方攻击时反弹 15 点直接伤害。",
        "cost": 3, "rarity": "rare", "type": "defense", "pool_type": "PVP", "upgradeable": 0, "level": 1,
        "effect_type": "shield", "damage": 0, "heal": 0, "shield": 20, "draw": 0, "energy": 0,
        "status": "", "buff": "thorns:15", "target": "self", "hit_count": 1, "self_damage": 0, "proc_chance": 0,
        "condition": "", "condition_param": 0, "condition_bonus": 0, "tags": "defense"
    },
    {
        "id": "pvp_feint", "name": "假动作", "description": "虚晃身形获得 10 点护盾，若本回合内受到攻击则立即触发应激反应追加获得 10 点护盾。",
        "cost": 1, "rarity": "common", "type": "defense", "pool_type": "PVP", "upgradeable": 0, "level": 1,
        "effect_type": "shield", "damage": 0, "heal": 0, "shield": 10, "draw": 0, "energy": 0,
        "status": "", "buff": "on_hit_shield:10", "target": "self", "hit_count": 1, "self_damage": 0, "proc_chance": 0,
        "condition": "", "condition_param": 0, "condition_bonus": 0, "tags": "defense"
    },
    {
        "id": "pvp_evasion", "name": "闪避", "description": "施展敏捷身法，使下一次受到的攻击伤害大幅降低 70%。",
        "cost": 2, "rarity": "rare", "type": "defense", "pool_type": "PVP", "upgradeable": 0, "level": 1,
        "effect_type": "buff", "damage": 0, "heal": 0, "shield": 0, "draw": 0, "energy": 0,
        "status": "", "buff": "next_damage_reduce:0.7", "target": "self", "hit_count": 1, "self_damage": 0, "proc_chance": 0,
        "condition": "", "condition_param": 0, "condition_bonus": 0, "tags": "defense"
    },
    {
        "id": "pvp_hold_formation", "name": "坚守阵型", "description": "本回合内受到的一切伤害降低 50%，但代价是本回合自身无法发动普通攻击。",
        "cost": 2, "rarity": "rare", "type": "defense", "pool_type": "PVP", "upgradeable": 0, "level": 1,
        "effect_type": "buff", "damage": 0, "heal": 0, "shield": 0, "draw": 0, "energy": 0,
        "status": "", "buff": "damage_reduce:0.5,forbid_attack", "target": "self", "hit_count": 1, "self_damage": 0, "proc_chance": 0,
        "condition": "", "condition_param": 0, "condition_bonus": 0, "tags": "defense"
    },
    {
        "id": "pvp_life_talisman", "name": "生命护符", "description": "激活远古守护符印，自身获得 20 点护盾并恢复 10 点生命值。",
        "cost": 3, "rarity": "rare", "type": "defense", "pool_type": "PVP", "upgradeable": 0, "level": 1,
        "effect_type": "shield", "damage": 0, "heal": 10, "shield": 20, "draw": 0, "energy": 0,
        "status": "", "buff": "", "target": "self", "hit_count": 1, "self_damage": 0, "proc_chance": 0,
        "condition": "", "condition_param": 0, "condition_bonus": 0, "tags": "defense,heal"
    },
    {
        "id": "pvp_last_will", "name": "最后意志", "description": "获得不屈意志庇佑：本回合若受到致命伤害，生命值强制保留 1 点，绝境不死！",
        "cost": 2, "rarity": "epic", "type": "defense", "pool_type": "PVP", "upgradeable": 0, "level": 1,
        "effect_type": "buff", "damage": 0, "heal": 0, "shield": 0, "draw": 0, "energy": 0,
        "status": "", "buff": "unyielding:1", "target": "self", "hit_count": 1, "self_damage": 0, "proc_chance": 0,
        "condition": "", "condition_param": 0, "condition_bonus": 0, "tags": "defense"
    },
    {
        "id": "pvp_damage_transfer", "name": "伤害转移", "description": "施展逆转结界：本回合受到的 50% 伤害将直接转化为生命恢复。",
        "cost": 3, "rarity": "epic", "type": "defense", "pool_type": "PVP", "upgradeable": 0, "level": 1,
        "effect_type": "buff", "damage": 0, "heal": 0, "shield": 0, "draw": 0, "energy": 0,
        "status": "", "buff": "damage_to_heal:0.5", "target": "self", "hit_count": 1, "self_damage": 0, "proc_chance": 0,
        "condition": "", "condition_param": 0, "condition_bonus": 0, "tags": "defense,heal"
    },
    {
        "id": "pvp_shield_charge", "name": "盾牌充能", "description": "消耗自身 10 点护盾转化为魔能，立即获得 3 点能量（若护盾不足10点则无法使用）。",
        "cost": 1, "rarity": "rare", "type": "defense", "pool_type": "PVP", "upgradeable": 0, "level": 1,
        "effect_type": "energy", "damage": 0, "heal": 0, "shield": -10, "draw": 0, "energy": 3,
        "status": "", "buff": "", "target": "self", "hit_count": 1, "self_damage": 0, "proc_chance": 0,
        "condition": "require_shield:10", "condition_param": 10, "condition_bonus": 0, "tags": "defense"
    },
    {
        "id": "pvp_defensive_counter", "name": "防御反制", "description": "进入截击架势：若敌方本回合发动攻击，其造成的伤害降低 60%。",
        "cost": 2, "rarity": "rare", "type": "defense", "pool_type": "PVP", "upgradeable": 0, "level": 1,
        "effect_type": "buff", "damage": 0, "heal": 0, "shield": 0, "draw": 0, "energy": 0,
        "status": "", "buff": "enemy_attack_reduce:0.6", "target": "self", "hit_count": 1, "self_damage": 0, "proc_chance": 0,
        "condition": "", "condition_param": 0, "condition_bonus": 0, "tags": "defense"
    },
    {
        "id": "pvp_divine_aegis", "name": "圣盾", "description": "降下至高圣洁神盾，立即获得 60 点巨额护盾（本局对战仅可生效使用一次）。",
        "cost": 5, "rarity": "epic", "type": "defense", "pool_type": "PVP", "upgradeable": 0, "level": 1,
        "effect_type": "shield", "damage": 0, "heal": 0, "shield": 60, "draw": 0, "energy": 0,
        "status": "", "buff": "limit_once_per_match", "target": "self", "hit_count": 1, "self_damage": 0, "proc_chance": 0,
        "condition": "", "condition_param": 0, "condition_bonus": 0, "tags": "defense"
    },

    # Resource (12)
    {
        "id": "pvp_greed", "name": "贪婪", "description": "消耗0费透支命运：立即抽取 2 张牌，但下回合开始时减少 2 点能量。",
        "cost": 0, "rarity": "rare", "type": "resource", "pool_type": "PVP", "upgradeable": 0, "level": 1,
        "effect_type": "draw", "damage": 0, "heal": 0, "shield": 0, "draw": 2, "energy": 0,
        "status": "", "buff": "next_turn_energy:-2", "target": "self", "hit_count": 1, "self_damage": 0, "proc_chance": 0,
        "condition": "", "condition_param": 0, "condition_bonus": 0, "tags": "resource"
    },
    {
        "id": "pvp_energy_storage", "name": "能量储存", "description": "将能量凝聚于符文阵列中：下回合开始时额外获得 3 点可用能量。",
        "cost": 1, "rarity": "common", "type": "resource", "pool_type": "PVP", "upgradeable": 0, "level": 1,
        "effect_type": "buff", "damage": 0, "heal": 0, "shield": 0, "draw": 0, "energy": 0,
        "status": "", "buff": "next_turn_energy:3", "target": "self", "hit_count": 1, "self_damage": 0, "proc_chance": 0,
        "condition": "", "condition_param": 0, "condition_bonus": 0, "tags": "resource"
    },
    {
        "id": "pvp_energy_steal", "name": "能量偷取", "description": "汲取敌方的源力：剥夺敌方 2 点能量，自身立即恢复 2 点能量。",
        "cost": 3, "rarity": "epic", "type": "resource", "pool_type": "PVP", "upgradeable": 0, "level": 1,
        "effect_type": "energy", "damage": 0, "heal": 0, "shield": 0, "draw": 0, "energy": 2,
        "status": "", "buff": "enemy_energy:-2", "target": "self", "hit_count": 1, "self_damage": 0, "proc_chance": 0,
        "condition": "", "condition_param": 0, "condition_bonus": 0, "tags": "resource"
    },
    {
        "id": "pvp_resource_convert", "name": "资源转换", "description": "弃置自身随机 1 张手牌，立即将其置换转化为 3 点能量。",
        "cost": 2, "rarity": "rare", "type": "resource", "pool_type": "PVP", "upgradeable": 0, "level": 1,
        "effect_type": "energy", "damage": 0, "heal": 0, "shield": 0, "draw": 0, "energy": 3,
        "status": "", "buff": "discard_self:1", "target": "self", "hit_count": 1, "self_damage": 0, "proc_chance": 0,
        "condition": "", "condition_param": 0, "condition_bonus": 0, "tags": "resource"
    },
    {
        "id": "pvp_insight", "name": "洞察", "description": "消耗1点能量，透视敌方手牌，随机揭示查看敌方 2 张手牌情报。",
        "cost": 1, "rarity": "common", "type": "resource", "pool_type": "PVP", "upgradeable": 0, "level": 1,
        "effect_type": "buff", "damage": 0, "heal": 0, "shield": 0, "draw": 0, "energy": 0,
        "status": "", "buff": "reveal_enemy_cards:2", "target": "self", "hit_count": 1, "self_damage": 0, "proc_chance": 0,
        "condition": "", "condition_param": 0, "condition_bonus": 0, "tags": "resource"
    },
    {
        "id": "pvp_duplicate", "name": "复制", "description": "解析镜像幻影，在自身手牌中完全复制自身随机 1 张手牌。",
        "cost": 4, "rarity": "epic", "type": "resource", "pool_type": "PVP", "upgradeable": 0, "level": 1,
        "effect_type": "buff", "damage": 0, "heal": 0, "shield": 0, "draw": 0, "energy": 0,
        "status": "", "buff": "clone_hand_card:1", "target": "self", "hit_count": 1, "self_damage": 0, "proc_chance": 0,
        "condition": "", "condition_param": 0, "condition_bonus": 0, "tags": "resource"
    },
    {
        "id": "pvp_card_swap", "name": "交换", "description": "扰乱因果律，双方各自随机交换 1 张手牌到对方手牌中。",
        "cost": 2, "rarity": "rare", "type": "resource", "pool_type": "PVP", "upgradeable": 0, "level": 1,
        "effect_type": "buff", "damage": 0, "heal": 0, "shield": 0, "draw": 0, "energy": 0,
        "status": "", "buff": "swap_random_card:1", "target": "self", "hit_count": 1, "self_damage": 0, "proc_chance": 0,
        "condition": "", "condition_param": 0, "condition_bonus": 0, "tags": "resource"
    },
    {
        "id": "pvp_destiny_gamble", "name": "命运赌博", "description": "抛起宿命金币：50% 概率获得 5 点能量，50% 概率扣除自身 15 点生命！",
        "cost": 0, "rarity": "rare", "type": "resource", "pool_type": "PVP", "upgradeable": 0, "level": 1,
        "effect_type": "buff", "damage": 0, "heal": 0, "shield": 0, "draw": 0, "energy": 0,
        "status": "", "buff": "gamble_energy_or_hp", "target": "self", "hit_count": 1, "self_damage": 0, "proc_chance": 0.5,
        "condition": "chance", "condition_param": 0.5, "condition_bonus": 0, "tags": "resource"
    },
    {
        "id": "pvp_suppression", "name": "压制", "description": "施展资源威压：若敌方当前手牌数量 > 5 张，则随机强制弃除其 1 张手牌。",
        "cost": 3, "rarity": "rare", "type": "resource", "pool_type": "PVP", "upgradeable": 0, "level": 1,
        "effect_type": "buff", "damage": 0, "heal": 0, "shield": 0, "draw": 0, "energy": 0,
        "status": "", "buff": "discard_enemy:1", "target": "enemy", "hit_count": 1, "self_damage": 0, "proc_chance": 0,
        "condition": "enemy_hand_gt_5", "condition_param": 5, "condition_bonus": 0, "tags": "resource"
    },
    {
        "id": "pvp_meditation", "name": "冥想", "description": "静心调息恢复 5 点生命，并顺势抽取 2 张手牌补充资源。",
        "cost": 1, "rarity": "common", "type": "resource", "pool_type": "PVP", "upgradeable": 0, "level": 1,
        "effect_type": "heal", "damage": 0, "heal": 5, "shield": 0, "draw": 2, "energy": 0,
        "status": "", "buff": "", "target": "self", "hit_count": 1, "self_damage": 0, "proc_chance": 0,
        "condition": "", "condition_param": 0, "condition_bonus": 0, "tags": "heal,resource"
    },
    {
        "id": "pvp_rapid_learning", "name": "快速学习", "description": "翻阅战术秘典，立即从牌库中抽取 3 张手牌。",
        "cost": 2, "rarity": "rare", "type": "resource", "pool_type": "PVP", "upgradeable": 0, "level": 1,
        "effect_type": "draw", "damage": 0, "heal": 0, "shield": 0, "draw": 3, "energy": 0,
        "status": "", "buff": "", "target": "self", "hit_count": 1, "self_damage": 0, "proc_chance": 0,
        "condition": "", "condition_param": 0, "condition_bonus": 0, "tags": "resource"
    },
    {
        "id": "pvp_energy_core", "name": "储能核心", "description": "植入永恒能量矩阵：使自身的能量上限永久 +1 点。",
        "cost": 3, "rarity": "epic", "type": "resource", "pool_type": "PVP", "upgradeable": 0, "level": 1,
        "effect_type": "buff", "damage": 0, "heal": 0, "shield": 0, "draw": 0, "energy": 0,
        "status": "", "buff": "max_energy+1", "target": "self", "hit_count": 1, "self_damage": 0, "proc_chance": 0,
        "condition": "", "condition_param": 0, "condition_bonus": 0, "tags": "resource"
    },

    # Control (12)
    {
        "id": "pvp_curse_of_weakness", "name": "虚弱诅咒", "description": "降下虚弱魔咒，使敌方造成的伤害削减 5 点，持续 2 回合。",
        "cost": 2, "rarity": "common", "type": "control", "pool_type": "PVP", "upgradeable": 0, "level": 1,
        "effect_type": "status", "damage": 0, "heal": 0, "shield": 0, "draw": 0, "energy": 0,
        "status": "weakness:5", "buff": "", "target": "enemy", "hit_count": 1, "self_damage": 0, "proc_chance": 0,
        "condition": "", "condition_param": 0, "condition_bonus": 0, "tags": "control"
    },
    {
        "id": "pvp_silence", "name": "沉默", "description": "释放禁魔力场，封印敌方英雄技能 1 回合，无法释放专属技能。",
        "cost": 4, "rarity": "epic", "type": "control", "pool_type": "PVP", "upgradeable": 0, "level": 1,
        "effect_type": "status", "damage": 0, "heal": 0, "shield": 0, "draw": 0, "energy": 0,
        "status": "silence:1", "buff": "", "target": "enemy", "hit_count": 1, "self_damage": 0, "proc_chance": 0,
        "condition": "", "condition_param": 0, "condition_bonus": 0, "tags": "control"
    },
    {
        "id": "pvp_deep_frost", "name": "冰封", "description": "席卷绝对寒潮，有 50% 概率深度冻结敌方 1 回合，使其下回合无法行动！",
        "cost": 3, "rarity": "rare", "type": "control", "pool_type": "PVP", "upgradeable": 0, "level": 1,
        "effect_type": "status", "damage": 0, "heal": 0, "shield": 0, "draw": 0, "energy": 0,
        "status": "freeze:1", "buff": "", "target": "enemy", "hit_count": 1, "self_damage": 0, "proc_chance": 0.5,
        "condition": "chance", "condition_param": 0.5, "condition_bonus": 0, "tags": "ice,control"
    },
    {
        "id": "pvp_mortal_strike", "name": "禁疗", "description": "重创敌方生机，使目标在接下来的回合内受到的所有治疗效果降低 50%。",
        "cost": 3, "rarity": "rare", "type": "control", "pool_type": "PVP", "upgradeable": 0, "level": 1,
        "effect_type": "status", "damage": 0, "heal": 0, "shield": 0, "draw": 0, "energy": 0,
        "status": "anti_heal:0.5", "buff": "", "target": "enemy", "hit_count": 1, "self_damage": 0, "proc_chance": 0,
        "condition": "", "condition_param": 0, "condition_bonus": 0, "tags": "control"
    },
    {
        "id": "pvp_dispel", "name": "驱散", "description": "释放奥术冲击，立即解除驱散敌方身上的所有正面增益状态（Buff）。",
        "cost": 2, "rarity": "rare", "type": "control", "pool_type": "PVP", "upgradeable": 0, "level": 1,
        "effect_type": "status", "damage": 0, "heal": 0, "shield": 0, "draw": 0, "energy": 0,
        "status": "remove_positive", "buff": "", "target": "enemy", "hit_count": 1, "self_damage": 0, "proc_chance": 0,
        "condition": "", "condition_param": 0, "condition_bonus": 0, "tags": "control"
    },
    {
        "id": "pvp_mark", "name": "标记", "description": "打上猎手血印，使目标受到的一切伤害增加 20%，持续 2 回合。",
        "cost": 1, "rarity": "common", "type": "control", "pool_type": "PVP", "upgradeable": 0, "level": 1,
        "effect_type": "status", "damage": 0, "heal": 0, "shield": 0, "draw": 0, "energy": 0,
        "status": "vulnerable:0.2", "buff": "", "target": "enemy", "hit_count": 1, "self_damage": 0, "proc_chance": 0,
        "condition": "", "condition_param": 0, "condition_bonus": 0, "tags": "control"
    },
    {
        "id": "pvp_chaos", "name": "混乱", "description": "扰乱敌方施法思绪，随机使敌方手牌中的 1 张卡牌能量费用 +2 点。",
        "cost": 2, "rarity": "rare", "type": "control", "pool_type": "PVP", "upgradeable": 0, "level": 1,
        "effect_type": "buff", "damage": 0, "heal": 0, "shield": 0, "draw": 0, "energy": 0,
        "status": "", "buff": "enemy_card_cost+2", "target": "enemy", "hit_count": 1, "self_damage": 0, "proc_chance": 0,
        "condition": "", "condition_param": 0, "condition_bonus": 0, "tags": "control"
    },
    {
        "id": "pvp_delay_curse", "name": "延迟诅咒", "description": "施加命运迟滞枷锁：敌方在下一个回合抽牌阶段少抽 2 张卡牌。",
        "cost": 2, "rarity": "rare", "type": "control", "pool_type": "PVP", "upgradeable": 0, "level": 1,
        "effect_type": "buff", "damage": 0, "heal": 0, "shield": 0, "draw": 0, "energy": 0,
        "status": "", "buff": "enemy_next_turn_draw:-2", "target": "enemy", "hit_count": 1, "self_damage": 0, "proc_chance": 0,
        "condition": "", "condition_param": 0, "condition_bonus": 0, "tags": "control"
    },
    {
        "id": "pvp_reversal", "name": "反转", "description": "力场扭曲：直接将敌方当前拥有的全部护盾值瞬间转化为直接伤害全额扣除！",
        "cost": 4, "rarity": "epic", "type": "control", "pool_type": "PVP", "upgradeable": 0, "level": 1,
        "effect_type": "buff", "damage": 0, "heal": 0, "shield": 0, "draw": 0, "energy": 0,
        "status": "", "buff": "shield_to_damage:1.0", "target": "enemy", "hit_count": 1, "self_damage": 0, "proc_chance": 0,
        "condition": "", "condition_param": 0, "condition_bonus": 0, "tags": "control"
    },
    {
        "id": "pvp_anti_magic_field", "name": "禁魔领域", "description": "展开禁魔结界：目标在接下来的 2 回合内完全无法获得任何护盾！",
        "cost": 5, "rarity": "epic", "type": "control", "pool_type": "PVP", "upgradeable": 0, "level": 1,
        "effect_type": "status", "damage": 0, "heal": 0, "shield": 0, "draw": 0, "energy": 0,
        "status": "no_shield:2", "buff": "", "target": "enemy", "hit_count": 1, "self_damage": 0, "proc_chance": 0,
        "condition": "", "condition_param": 0, "condition_bonus": 0, "tags": "control"
    },
    {
        "id": "pvp_time_stasis", "name": "时间暂停", "description": "撕裂时间缝隙，强制剥夺限制敌方下回合所有行动机会！",
        "cost": 5, "rarity": "epic", "type": "control", "pool_type": "PVP", "upgradeable": 0, "level": 1,
        "effect_type": "status", "damage": 0, "heal": 0, "shield": 0, "draw": 0, "energy": 0,
        "status": "stasis:1", "buff": "", "target": "enemy", "hit_count": 1, "self_damage": 0, "proc_chance": 0,
        "condition": "", "condition_param": 0, "condition_bonus": 0, "tags": "control"
    },
    {
        "id": "pvp_purifying_field", "name": "净化领域", "description": "激发神圣净化结界：驱散解除战场上双方所有的持续伤害状态（灼烧、剧毒等）。",
        "cost": 3, "rarity": "rare", "type": "control", "pool_type": "PVP", "upgradeable": 0, "level": 1,
        "effect_type": "status", "damage": 0, "heal": 0, "shield": 0, "draw": 0, "energy": 0,
        "status": "dispel_all_dots", "buff": "", "target": "self", "hit_count": 1, "self_damage": 0, "proc_chance": 0,
        "condition": "", "condition_param": 0, "condition_bonus": 0, "tags": "control"
    }
]

# 3. PVE Cards (48: 6 Archetypes x 8 Cards, upgradeable=true, level=1-3)
# Fire(8), Ice(8), Poison(8), Shield(8), Heal(8), Energy(8)
PVE_CARDS = [
    # 1. Fire (8)
    {
        "id": "pve_flame_spark", "name": "火花印记", "description": "造成 7 点火焰伤害，并附加 1 层灼烧。",
        "cost": 1, "rarity": "common", "type": "attack", "pool_type": "PVE", "upgradeable": 1, "level": 1,
        "effect_type": "damage", "damage": 7, "heal": 0, "shield": 0, "draw": 0, "energy": 0,
        "status": "burn:1", "buff": "", "target": "enemy", "hit_count": 1, "self_damage": 0, "proc_chance": 0,
        "condition": "", "condition_param": 0, "condition_bonus": 0, "tags": "fire,attack"
    },
    {
        "id": "pve_ignite", "name": "点燃", "description": "引燃敌方核心造成 14 点伤害，并附加 2 层灼烧。",
        "cost": 2, "rarity": "rare", "type": "attack", "pool_type": "PVE", "upgradeable": 1, "level": 1,
        "effect_type": "damage", "damage": 14, "heal": 0, "shield": 0, "draw": 0, "energy": 0,
        "status": "burn:2", "buff": "", "target": "enemy", "hit_count": 1, "self_damage": 0, "proc_chance": 0,
        "condition": "", "condition_param": 0, "condition_bonus": 0, "tags": "fire,attack"
    },
    {
        "id": "pve_fireball", "name": "烈焰火球", "description": "发射巨大炽热火球造成 24 点伤害，并附加 3 层灼烧。",
        "cost": 3, "rarity": "rare", "type": "attack", "pool_type": "PVE", "upgradeable": 1, "level": 1,
        "effect_type": "damage", "damage": 24, "heal": 0, "shield": 0, "draw": 0, "energy": 0,
        "status": "burn:3", "buff": "", "target": "enemy", "hit_count": 1, "self_damage": 0, "proc_chance": 0,
        "condition": "", "condition_param": 0, "condition_bonus": 0, "tags": "fire,attack"
    },
    {
        "id": "pve_flame_barrier", "name": "炽焰屏障", "description": "烈火化为盾牌获得 16 点护盾，受击时反灼伤攻击者 1 层灼烧。",
        "cost": 2, "rarity": "rare", "type": "defense", "pool_type": "PVE", "upgradeable": 1, "level": 1,
        "effect_type": "shield", "damage": 0, "heal": 0, "shield": 16, "draw": 0, "energy": 0,
        "status": "", "buff": "burn_attacker:1", "target": "self", "hit_count": 1, "self_damage": 0, "proc_chance": 0,
        "condition": "", "condition_param": 0, "condition_bonus": 0, "tags": "fire,defense"
    },
    {
        "id": "pve_combustion", "name": "连环引燃", "description": "造成 12 点伤害。若目标处于灼烧状态，则触发引爆追加 10 点伤害。",
        "cost": 2, "rarity": "rare", "type": "attack", "pool_type": "PVE", "upgradeable": 1, "level": 1,
        "effect_type": "damage", "damage": 12, "heal": 0, "shield": 0, "draw": 0, "energy": 0,
        "status": "", "buff": "", "target": "enemy", "hit_count": 1, "self_damage": 0, "proc_chance": 0,
        "condition": "target_has_burn", "condition_param": 0, "condition_bonus": 10, "tags": "fire,attack"
    },
    {
        "id": "pve_pyroblast", "name": "炎爆术", "description": "引导强力炎爆对目标造成 32 点毁灭伤害，并附加 4 层灼烧。",
        "cost": 4, "rarity": "epic", "type": "attack", "pool_type": "PVE", "upgradeable": 1, "level": 1,
        "effect_type": "damage", "damage": 32, "heal": 0, "shield": 0, "draw": 0, "energy": 0,
        "status": "burn:4", "buff": "", "target": "enemy", "hit_count": 1, "self_damage": 0, "proc_chance": 0,
        "condition": "", "condition_param": 0, "condition_bonus": 0, "tags": "fire,attack"
    },
    {
        "id": "pve_fire_infusion", "name": "烈焰附魔", "description": "使武器与法术附带火灵共鸣，本回合所有火焰卡牌伤害 +6 点。",
        "cost": 1, "rarity": "rare", "type": "buff", "pool_type": "PVE", "upgradeable": 1, "level": 1,
        "effect_type": "buff", "damage": 0, "heal": 0, "shield": 0, "draw": 0, "energy": 0,
        "status": "", "buff": "fire_damage_bonus:6", "target": "self", "hit_count": 1, "self_damage": 0, "proc_chance": 0,
        "condition": "", "condition_param": 0, "condition_bonus": 0, "tags": "fire"
    },
    {
        "id": "pve_hellfire", "name": "地狱烈火", "description": "释放地狱硫磺烈焰造成 40 点伤害并附加 5 层灼烧，自身承受 6 点反冲伤害。",
        "cost": 5, "rarity": "epic", "type": "attack", "pool_type": "PVE", "upgradeable": 1, "level": 1,
        "effect_type": "damage", "damage": 40, "heal": 0, "shield": 0, "draw": 0, "energy": 0,
        "status": "burn:5", "buff": "", "target": "enemy", "hit_count": 1, "self_damage": 6, "proc_chance": 0,
        "condition": "", "condition_param": 0, "condition_bonus": 0, "tags": "fire,attack"
    },

    # 2. Ice (8)
    {
        "id": "pve_ice_shard", "name": "冰凌碎片", "description": "发射轻盈冰棱造成 6 点伤害，附加 1 层霜冻减速。",
        "cost": 1, "rarity": "common", "type": "attack", "pool_type": "PVE", "upgradeable": 1, "level": 1,
        "effect_type": "damage", "damage": 6, "heal": 0, "shield": 0, "draw": 0, "energy": 0,
        "status": "chill:1", "buff": "", "target": "enemy", "hit_count": 1, "self_damage": 0, "proc_chance": 0,
        "condition": "", "condition_param": 0, "condition_bonus": 0, "tags": "ice,attack"
    },
    {
        "id": "pve_frost_lance", "name": "寒霜刺枪", "description": "投掷凝冰长枪造成 13 点伤害。若目标处于冻结或霜冻状态，追加 12 点破冰伤害。",
        "cost": 2, "rarity": "rare", "type": "attack", "pool_type": "PVE", "upgradeable": 1, "level": 1,
        "effect_type": "damage", "damage": 13, "heal": 0, "shield": 0, "draw": 0, "energy": 0,
        "status": "", "buff": "", "target": "enemy", "hit_count": 1, "self_damage": 0, "proc_chance": 0,
        "condition": "target_frozen", "condition_param": 0, "condition_bonus": 12, "tags": "ice,attack"
    },
    {
        "id": "pve_blizzard", "name": "暴风雪", "description": "召唤冰风暴造成 18 点伤害，并有 40% 几率直接冻结目标 1 回合。",
        "cost": 3, "rarity": "rare", "type": "attack", "pool_type": "PVE", "upgradeable": 1, "level": 1,
        "effect_type": "damage", "damage": 18, "heal": 0, "shield": 0, "draw": 0, "energy": 0,
        "status": "freeze:1", "buff": "", "target": "enemy", "hit_count": 1, "self_damage": 0, "proc_chance": 0.4,
        "condition": "chance", "condition_param": 0.4, "condition_bonus": 0, "tags": "ice,attack"
    },
    {
        "id": "pve_ice_armor", "name": "寒冰重铠", "description": "以万年玄冰凝结重铠获得 20 点护盾，受击时减速攻击者 1 层霜冻。",
        "cost": 2, "rarity": "rare", "type": "defense", "pool_type": "PVE", "upgradeable": 1, "level": 1,
        "effect_type": "shield", "damage": 0, "heal": 0, "shield": 20, "draw": 0, "energy": 0,
        "status": "", "buff": "chill_attacker:1", "target": "self", "hit_count": 1, "self_damage": 0, "proc_chance": 0,
        "condition": "", "condition_param": 0, "condition_bonus": 0, "tags": "ice,defense"
    },
    {
        "id": "pve_glacial_shield", "name": "极寒坚冰", "description": "筑起绝对零度冰壁获得 32 点护盾，并有 30% 概率将攻击者完全冻结。",
        "cost": 3, "rarity": "rare", "type": "defense", "pool_type": "PVE", "upgradeable": 1, "level": 1,
        "effect_type": "shield", "damage": 0, "heal": 0, "shield": 32, "draw": 0, "energy": 0,
        "status": "freeze:1", "buff": "", "target": "self", "hit_count": 1, "self_damage": 0, "proc_chance": 0.3,
        "condition": "chance", "condition_param": 0.3, "condition_bonus": 0, "tags": "ice,defense"
    },
    {
        "id": "pve_frost_nova", "name": "冰霜新星", "description": "在目标脚下引爆冰环造成 15 点伤害，100% 深度冻结目标 1 回合。",
        "cost": 3, "rarity": "epic", "type": "attack", "pool_type": "PVE", "upgradeable": 1, "level": 1,
        "effect_type": "damage", "damage": 15, "heal": 0, "shield": 0, "draw": 0, "energy": 0,
        "status": "freeze:1", "buff": "", "target": "enemy", "hit_count": 1, "self_damage": 0, "proc_chance": 1.0,
        "condition": "", "condition_param": 0, "condition_bonus": 0, "tags": "ice,attack"
    },
    {
        "id": "pve_absolute_zero", "name": "绝对零度", "description": "释放极冻死光造成 26 点极寒伤害，并冻结敌方 1 回合。",
        "cost": 4, "rarity": "epic", "type": "attack", "pool_type": "PVE", "upgradeable": 1, "level": 1,
        "effect_type": "damage", "damage": 26, "heal": 0, "shield": 0, "draw": 0, "energy": 0,
        "status": "freeze:1", "buff": "", "target": "enemy", "hit_count": 1, "self_damage": 0, "proc_chance": 1.0,
        "condition": "", "condition_param": 0, "condition_bonus": 0, "tags": "ice,attack"
    },
    {
        "id": "pve_frozen_orb", "name": "寒冰宝珠", "description": "掷出旋转冰珠造成 10 点穿透伤害，并在冰雾中抽取 2 张手牌。",
        "cost": 2, "rarity": "rare", "type": "attack", "pool_type": "PVE", "upgradeable": 1, "level": 1,
        "effect_type": "damage", "damage": 10, "heal": 0, "shield": 0, "draw": 2, "energy": 0,
        "status": "", "buff": "", "target": "enemy", "hit_count": 1, "self_damage": 0, "proc_chance": 0,
        "condition": "", "condition_param": 0, "condition_bonus": 0, "tags": "ice,attack"
    },

    # 3. Poison (8)
    {
        "id": "pve_toxic_dart", "name": "淬毒飞镖", "description": "射出暗影毒镖造成 5 点伤害，并附加 2 层剧毒状态。",
        "cost": 1, "rarity": "common", "type": "attack", "pool_type": "PVE", "upgradeable": 1, "level": 1,
        "effect_type": "damage", "damage": 5, "heal": 0, "shield": 0, "draw": 0, "energy": 0,
        "status": "poison:2", "buff": "", "target": "enemy", "hit_count": 1, "self_damage": 0, "proc_chance": 0,
        "condition": "", "condition_param": 0, "condition_bonus": 0, "tags": "poison,attack"
    },
    {
        "id": "pve_venom_slash", "name": "剧毒之刺", "description": "以剧毒弯刀深刺目标造成 11 点伤害，附加 4 层剧毒。",
        "cost": 2, "rarity": "rare", "type": "attack", "pool_type": "PVE", "upgradeable": 1, "level": 1,
        "effect_type": "damage", "damage": 11, "heal": 0, "shield": 0, "draw": 0, "energy": 0,
        "status": "poison:4", "buff": "", "target": "enemy", "hit_count": 1, "self_damage": 0, "proc_chance": 0,
        "condition": "", "condition_param": 0, "condition_bonus": 0, "tags": "poison,attack"
    },
    {
        "id": "pve_poison_mist", "name": "毒云弥漫", "description": "释放大范围腐蚀毒雾，直接为目标叠加大规模 7 层剧毒！",
        "cost": 3, "rarity": "rare", "type": "status", "pool_type": "PVE", "upgradeable": 1, "level": 1,
        "effect_type": "status", "damage": 0, "heal": 0, "shield": 0, "draw": 0, "energy": 0,
        "status": "poison:7", "buff": "", "target": "enemy", "hit_count": 1, "self_damage": 0, "proc_chance": 0,
        "condition": "", "condition_param": 0, "condition_bonus": 0, "tags": "poison"
    },
    {
        "id": "pve_acid_shield", "name": "强酸护盾", "description": "凝聚强酸能量护盾获得 15 点护盾，受击时反噬攻击者 2 层剧毒。",
        "cost": 2, "rarity": "rare", "type": "defense", "pool_type": "PVE", "upgradeable": 1, "level": 1,
        "effect_type": "shield", "damage": 0, "heal": 0, "shield": 15, "draw": 0, "energy": 0,
        "status": "", "buff": "poison_attacker:2", "target": "self", "hit_count": 1, "self_damage": 0, "proc_chance": 0,
        "condition": "", "condition_param": 0, "condition_bonus": 0, "tags": "poison,defense"
    },
    {
        "id": "pve_catalyst", "name": "剧毒催化", "description": "造成 10 点毒性震荡伤害。若目标身上有剧毒，则催化猛烈追加 15 点爆发伤害！",
        "cost": 2, "rarity": "rare", "type": "attack", "pool_type": "PVE", "upgradeable": 1, "level": 1,
        "effect_type": "damage", "damage": 10, "heal": 0, "shield": 0, "draw": 0, "energy": 0,
        "status": "", "buff": "", "target": "enemy", "hit_count": 1, "self_damage": 0, "proc_chance": 0,
        "condition": "target_has_poison", "condition_param": 0, "condition_bonus": 15, "tags": "poison,attack"
    },
    {
        "id": "pve_neurotoxin", "name": "神经毒素", "description": "注入破坏神经的剧毒造成 14 点伤害，附加 3 层剧毒并削弱目标 3 点伤害。",
        "cost": 3, "rarity": "rare", "type": "attack", "pool_type": "PVE", "upgradeable": 1, "level": 1,
        "effect_type": "damage", "damage": 14, "heal": 0, "shield": 0, "draw": 0, "energy": 0,
        "status": "poison:3", "buff": "weaken:3", "target": "enemy", "hit_count": 1, "self_damage": 0, "proc_chance": 0,
        "condition": "", "condition_param": 0, "condition_bonus": 0, "tags": "poison,attack"
    },
    {
        "id": "pve_plague_strike", "name": "瘟疫打击", "description": "汇聚瘟疫重击造成 22 点伤害，并附加 6 层致命剧毒。",
        "cost": 4, "rarity": "epic", "type": "attack", "pool_type": "PVE", "upgradeable": 1, "level": 1,
        "effect_type": "damage", "damage": 22, "heal": 0, "shield": 0, "draw": 0, "energy": 0,
        "status": "poison:6", "buff": "", "target": "enemy", "hit_count": 1, "self_damage": 0, "proc_chance": 0,
        "condition": "", "condition_param": 0, "condition_bonus": 0, "tags": "poison,attack"
    },
    {
        "id": "pve_toxic_infusion", "name": "百毒附体", "description": "吞服百草毒丹：本回合所有剧毒卡牌造成的伤害和层数 +5 点。",
        "cost": 1, "rarity": "rare", "type": "buff", "pool_type": "PVE", "upgradeable": 1, "level": 1,
        "effect_type": "buff", "damage": 0, "heal": 0, "shield": 0, "draw": 0, "energy": 0,
        "status": "", "buff": "poison_bonus:5", "target": "self", "hit_count": 1, "self_damage": 0, "proc_chance": 0,
        "condition": "", "condition_param": 0, "condition_bonus": 0, "tags": "poison"
    },

    # 4. Shield (8)
    {
        "id": "pve_buckler", "name": "强化小圆盾", "description": "熟练格挡，自身立即获得 14 点护盾。",
        "cost": 1, "rarity": "common", "type": "defense", "pool_type": "PVE", "upgradeable": 1, "level": 1,
        "effect_type": "shield", "damage": 0, "heal": 0, "shield": 14, "draw": 0, "energy": 0,
        "status": "", "buff": "", "target": "self", "hit_count": 1, "self_damage": 0, "proc_chance": 0,
        "condition": "", "condition_param": 0, "condition_bonus": 0, "tags": "defense"
    },
    {
        "id": "pve_shield_slam", "name": "盾墙猛击", "description": "持盾向前撞击造成 15 点物理伤害。若自身持有护盾，额外追加 10 点重击伤害。",
        "cost": 2, "rarity": "rare", "type": "attack", "pool_type": "PVE", "upgradeable": 1, "level": 1,
        "effect_type": "damage", "damage": 15, "heal": 0, "shield": 0, "draw": 0, "energy": 0,
        "status": "", "buff": "", "target": "enemy", "hit_count": 1, "self_damage": 0, "proc_chance": 0,
        "condition": "user_has_shield", "condition_param": 0, "condition_bonus": 10, "tags": "defense,attack"
    },
    {
        "id": "pve_solid_barrier", "name": "坚实壁障", "description": "构筑高耸岩石护盾，自身立即获得 26 点护盾。",
        "cost": 2, "rarity": "rare", "type": "defense", "pool_type": "PVE", "upgradeable": 1, "level": 1,
        "effect_type": "shield", "damage": 0, "heal": 0, "shield": 26, "draw": 0, "energy": 0,
        "status": "", "buff": "", "target": "self", "hit_count": 1, "self_damage": 0, "proc_chance": 0,
        "condition": "", "condition_param": 0, "condition_bonus": 0, "tags": "defense"
    },
    {
        "id": "pve_bastion", "name": "叹息之壁", "description": "唤起不可逾越之绝壁，自身获得 42 点巨额护盾。",
        "cost": 3, "rarity": "rare", "type": "defense", "pool_type": "PVE", "upgradeable": 1, "level": 1,
        "effect_type": "shield", "damage": 0, "heal": 0, "shield": 42, "draw": 0, "energy": 0,
        "status": "", "buff": "", "target": "self", "hit_count": 1, "self_damage": 0, "proc_chance": 0,
        "condition": "", "condition_param": 0, "condition_bonus": 0, "tags": "defense"
    },
    {
        "id": "pve_iron_will", "name": "钢铁意志", "description": "心如精钢，本回合自身获得的所有护盾效果提升 +8 点。",
        "cost": 1, "rarity": "rare", "type": "buff", "pool_type": "PVE", "upgradeable": 1, "level": 1,
        "effect_type": "buff", "damage": 0, "heal": 0, "shield": 0, "draw": 0, "energy": 0,
        "status": "", "buff": "shield_bonus:8", "target": "self", "hit_count": 1, "self_damage": 0, "proc_chance": 0,
        "condition": "", "condition_param": 0, "condition_bonus": 0, "tags": "defense"
    },
    {
        "id": "pve_shield_burst", "name": "护盾迸裂", "description": "攻防兼备：自身获得 16 点护盾，同时向前迸发能量冲击造成 16 点伤害。",
        "cost": 3, "rarity": "rare", "type": "attack", "pool_type": "PVE", "upgradeable": 1, "level": 1,
        "effect_type": "damage", "damage": 16, "heal": 0, "shield": 16, "draw": 0, "energy": 0,
        "status": "", "buff": "", "target": "enemy", "hit_count": 1, "self_damage": 0, "proc_chance": 0,
        "condition": "", "condition_param": 0, "condition_bonus": 0, "tags": "defense,attack"
    },
    {
        "id": "pve_guardian_aura", "name": "守护光环", "description": "释放庇护光环获得 20 点护盾，本回合内受到的所有单次伤害额外固定减少 3 点。",
        "cost": 2, "rarity": "epic", "type": "defense", "pool_type": "PVE", "upgradeable": 1, "level": 1,
        "effect_type": "shield", "damage": 0, "heal": 0, "shield": 20, "draw": 0, "energy": 0,
        "status": "", "buff": "flat_reduction:3", "target": "self", "hit_count": 1, "self_damage": 0, "proc_chance": 0,
        "condition": "", "condition_param": 0, "condition_bonus": 0, "tags": "defense"
    },
    {
        "id": "pve_fortress", "name": "移动要塞", "description": "彻底化身坚固战阵堡垒，自身立即获得 55 点护盾！",
        "cost": 4, "rarity": "epic", "type": "defense", "pool_type": "PVE", "upgradeable": 1, "level": 1,
        "effect_type": "shield", "damage": 0, "heal": 0, "shield": 55, "draw": 0, "energy": 0,
        "status": "", "buff": "", "target": "self", "hit_count": 1, "self_damage": 0, "proc_chance": 0,
        "condition": "", "condition_param": 0, "condition_bonus": 0, "tags": "defense"
    },

    # 5. Heal (8)
    {
        "id": "pve_minor_mending", "name": "次级愈合", "description": "引导微光愈合伤势，自身恢复 10 点生命值。",
        "cost": 1, "rarity": "common", "type": "heal", "pool_type": "PVE", "upgradeable": 1, "level": 1,
        "effect_type": "heal", "damage": 0, "heal": 10, "shield": 0, "draw": 0, "energy": 0,
        "status": "", "buff": "", "target": "self", "hit_count": 1, "self_damage": 0, "proc_chance": 0,
        "condition": "", "condition_param": 0, "condition_bonus": 0, "tags": "heal"
    },
    {
        "id": "pve_holy_light", "name": "圣光闪耀", "description": "沐浴温暖圣光，自身恢复 20 点生命值。",
        "cost": 2, "rarity": "rare", "type": "heal", "pool_type": "PVE", "upgradeable": 1, "level": 1,
        "effect_type": "heal", "damage": 0, "heal": 20, "shield": 0, "draw": 0, "energy": 0,
        "status": "", "buff": "", "target": "self", "hit_count": 1, "self_damage": 0, "proc_chance": 0,
        "condition": "", "condition_param": 0, "condition_bonus": 0, "tags": "heal"
    },
    {
        "id": "pve_divine_blessing", "name": "神圣恩赐", "description": "祈愿神圣复苏甘霖，自身恢复 32 点生命值。",
        "cost": 3, "rarity": "rare", "type": "heal", "pool_type": "PVE", "upgradeable": 1, "level": 1,
        "effect_type": "heal", "damage": 0, "heal": 32, "shield": 0, "draw": 0, "energy": 0,
        "status": "", "buff": "", "target": "self", "hit_count": 1, "self_damage": 0, "proc_chance": 0,
        "condition": "", "condition_param": 0, "condition_bonus": 0, "tags": "heal"
    },
    {
        "id": "pve_regeneration", "name": "自然复苏", "description": "汲取自然生机立即恢复 12 点生命，并在回合结束时持续自愈 5 点生命。",
        "cost": 2, "rarity": "rare", "type": "heal", "pool_type": "PVE", "upgradeable": 1, "level": 1,
        "effect_type": "heal", "damage": 0, "heal": 12, "shield": 0, "draw": 0, "energy": 0,
        "status": "", "buff": "turn_end_heal:5", "target": "self", "hit_count": 1, "self_damage": 0, "proc_chance": 0,
        "condition": "", "condition_param": 0, "condition_bonus": 0, "tags": "heal"
    },
    {
        "id": "pve_healing_ward", "name": "治疗图腾", "description": "召唤守护图腾，同时为自身恢复 15 点生命并获得 15 点护盾。",
        "cost": 2, "rarity": "rare", "type": "heal", "pool_type": "PVE", "upgradeable": 1, "level": 1,
        "effect_type": "heal", "damage": 0, "heal": 15, "shield": 15, "draw": 0, "energy": 0,
        "status": "", "buff": "", "target": "self", "hit_count": 1, "self_damage": 0, "proc_chance": 0,
        "condition": "", "condition_param": 0, "condition_bonus": 0, "tags": "heal,defense"
    },
    {
        "id": "pve_smite", "name": "惩戒圣击", "description": "圣光化作利刃对敌人造成 18 点伤害，同时反哺自身恢复 12 点生命值。",
        "cost": 3, "rarity": "rare", "type": "attack", "pool_type": "PVE", "upgradeable": 1, "level": 1,
        "effect_type": "damage", "damage": 18, "heal": 12, "shield": 0, "draw": 0, "energy": 0,
        "status": "", "buff": "", "target": "enemy", "hit_count": 1, "self_damage": 0, "proc_chance": 0,
        "condition": "", "condition_param": 0, "condition_bonus": 0, "tags": "heal,attack"
    },
    {
        "id": "pve_vitality_infusion", "name": "生命灌注", "description": "唤醒体内生命源泉：本回合内自身受到的所有治疗效果提升 +8 点。",
        "cost": 1, "rarity": "rare", "type": "buff", "pool_type": "PVE", "upgradeable": 1, "level": 1,
        "effect_type": "buff", "damage": 0, "heal": 0, "shield": 0, "draw": 0, "energy": 0,
        "status": "", "buff": "heal_bonus:8", "target": "self", "hit_count": 1, "self_damage": 0, "proc_chance": 0,
        "condition": "", "condition_param": 0, "condition_bonus": 0, "tags": "heal"
    },
    {
        "id": "pve_miracle", "name": "奇迹再临", "description": "神圣大奇迹降临，令濒死者重生：立即为自身恢复 45 点巨额生命值！",
        "cost": 4, "rarity": "epic", "type": "heal", "pool_type": "PVE", "upgradeable": 1, "level": 1,
        "effect_type": "heal", "damage": 0, "heal": 45, "shield": 0, "draw": 0, "energy": 0,
        "status": "", "buff": "", "target": "self", "hit_count": 1, "self_damage": 0, "proc_chance": 0,
        "condition": "", "condition_param": 0, "condition_bonus": 0, "tags": "heal"
    },

    # 6. Energy (8)
    {
        "id": "pve_spark", "name": "微光引导", "description": "消耗0点能量，引导自然灵气聚集，自身立即获得 2 点能量。",
        "cost": 0, "rarity": "common", "type": "resource", "pool_type": "PVE", "upgradeable": 1, "level": 1,
        "effect_type": "energy", "damage": 0, "heal": 0, "shield": 0, "draw": 0, "energy": 2,
        "status": "", "buff": "", "target": "self", "hit_count": 1, "self_damage": 0, "proc_chance": 0,
        "condition": "", "condition_param": 0, "condition_bonus": 0, "tags": "resource"
    },
    {
        "id": "pve_mana_channel", "name": "法力导流", "description": "疏通法力回路，自身获得 3 点能量并抽取 1 张手牌。",
        "cost": 1, "rarity": "rare", "type": "resource", "pool_type": "PVE", "upgradeable": 1, "level": 1,
        "effect_type": "energy", "damage": 0, "heal": 0, "shield": 0, "draw": 1, "energy": 3,
        "status": "", "buff": "", "target": "self", "hit_count": 1, "self_damage": 0, "proc_chance": 0,
        "condition": "", "condition_param": 0, "condition_bonus": 0, "tags": "resource"
    },
    {
        "id": "pve_power_infusion", "name": "能量充能", "description": "汇聚奥术结晶，自身立即充能获得 4 点可用能量。",
        "cost": 2, "rarity": "rare", "type": "resource", "pool_type": "PVE", "upgradeable": 1, "level": 1,
        "effect_type": "energy", "damage": 0, "heal": 0, "shield": 0, "draw": 0, "energy": 4,
        "status": "", "buff": "", "target": "self", "hit_count": 1, "self_damage": 0, "proc_chance": 0,
        "condition": "", "condition_param": 0, "condition_bonus": 0, "tags": "resource"
    },
    {
        "id": "pve_overload_bolt", "name": "过载电弧", "description": "释放高压电弧轰击目标造成 24 点雷电伤害，并在电磁残留中回充 1 点能量。",
        "cost": 3, "rarity": "rare", "type": "attack", "pool_type": "PVE", "upgradeable": 1, "level": 1,
        "effect_type": "damage", "damage": 24, "heal": 0, "shield": 0, "draw": 0, "energy": 1,
        "status": "", "buff": "", "target": "enemy", "hit_count": 1, "self_damage": 0, "proc_chance": 0,
        "condition": "", "condition_param": 0, "condition_bonus": 0, "tags": "attack,resource"
    },
    {
        "id": "pve_energy_barrier", "name": "法力护盾", "description": "将魔能转化为防御力场获得 20 点护盾，同时回流 1 点能量。",
        "cost": 1, "rarity": "rare", "type": "defense", "pool_type": "PVE", "upgradeable": 1, "level": 1,
        "effect_type": "shield", "damage": 0, "heal": 0, "shield": 20, "draw": 0, "energy": 1,
        "status": "", "buff": "", "target": "self", "hit_count": 1, "self_damage": 0, "proc_chance": 0,
        "condition": "", "condition_param": 0, "condition_bonus": 0, "tags": "defense,resource"
    },
    {
        "id": "pve_mana_overflow", "name": "能量溢出", "description": "魔力潮汐汹涌爆发：立即抽取 3 张牌，同时返还 2 点能量。",
        "cost": 2, "rarity": "epic", "type": "resource", "pool_type": "PVE", "upgradeable": 1, "level": 1,
        "effect_type": "draw", "damage": 0, "heal": 0, "shield": 0, "draw": 3, "energy": 2,
        "status": "", "buff": "", "target": "self", "hit_count": 1, "self_damage": 0, "proc_chance": 0,
        "condition": "", "condition_param": 0, "condition_bonus": 0, "tags": "resource"
    },
    {
        "id": "pve_arcane_storm", "name": "奥术风暴", "description": "引爆暴烈奥术风暴对目标造成 30 点魔法伤害，并回馈 1 点能量。",
        "cost": 4, "rarity": "epic", "type": "attack", "pool_type": "PVE", "upgradeable": 1, "level": 1,
        "effect_type": "damage", "damage": 30, "heal": 0, "shield": 0, "draw": 0, "energy": 1,
        "status": "", "buff": "", "target": "enemy", "hit_count": 1, "self_damage": 0, "proc_chance": 0,
        "condition": "", "condition_param": 0, "condition_bonus": 0, "tags": "attack,resource"
    },
    {
        "id": "pve_singularity", "name": "源质奇点", "description": "开启源质能量黑洞：直接充能 6 点能量，并大量抽取 3 张手牌！",
        "cost": 5, "rarity": "epic", "type": "resource", "pool_type": "PVE", "upgradeable": 1, "level": 1,
        "effect_type": "energy", "damage": 0, "heal": 0, "shield": 0, "draw": 3, "energy": 6,
        "status": "", "buff": "", "target": "self", "hit_count": 1, "self_damage": 0, "proc_chance": 0,
        "condition": "", "condition_param": 0, "condition_bonus": 0, "tags": "resource"
    }
]

# 4. Effect Dictionary
EFFECT_DICT = [
    {"effect_type": "damage", "name": "伤害效果", "description": "通过DamageCalculator执行标准流水线伤害计算，支持连击数(hit_count)、反噬伤害(self_damage)、概率追伤(proc_chance)及触发条件(condition)。", "target": "enemy"},
    {"effect_type": "heal", "name": "治疗效果", "description": "恢复目标或自身的生命值，享受healBonus治疗加成修正，受禁疗(anti_heal)状态削减。", "target": "self/enemy"},
    {"effect_type": "shield", "name": "护盾效果", "description": "为目标或自身提供护盾吸收伤害，享受shieldBonus护盾加成修正，受禁魔领域(no_shield)状态限制。", "target": "self/enemy"},
    {"effect_type": "draw", "name": "抽牌效果", "description": "从玩家的摸牌堆中抽取指定数量的手牌至手牌区，不超过手牌上限(hand_card_limit)。", "target": "self"},
    {"effect_type": "energy", "name": "能量效果", "description": "使目标或自身的当前能量增加或减少，上限受max_energy规则限制。", "target": "self/enemy"},
    {"effect_type": "status", "name": "状态效果", "description": "通过StatusSystem挂载状态（如灼烧burn、剧毒poison、冰冻freeze、沉默silence等），或解除异常状态(remove_negative)。", "target": "self/enemy"},
    {"effect_type": "buff", "name": "增益效果", "description": "赋予战斗增益或被动响应（如专注加伤damage+N、受击减伤damage_reduce、反伤thorns、不屈最后意志unyielding等）。", "target": "self/enemy"}
]

def add_sheet_data(ws, title, card_list):
    ws.title = title
    
    # Add headers
    for col_idx, (col_key, col_label, col_width) in enumerate(COLUMNS, 1):
        cell = ws.cell(row=1, column=col_idx, value=col_label)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = ALIGN_CENTER
        ws.column_dimensions[get_column_letter(col_idx)].width = col_width

    # Add rows
    for row_idx, item in enumerate(card_list, 2):
        for col_idx, (col_key, col_label, _) in enumerate(COLUMNS, 1):
            val = item.get(col_key, "")
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.font = CELL_FONT
            cell.border = BORDER_THIN
            if isinstance(val, (int, float)):
                cell.alignment = ALIGN_CENTER
            elif col_key in ("id", "rarity", "type", "pool_type", "effect_type", "target", "upgradeable", "level"):
                cell.alignment = ALIGN_CENTER
            else:
                cell.alignment = ALIGN_LEFT

def create_effects_sheet(ws):
    ws.title = "Card_Effects"
    headers = [("effect_type", "效果类型标识", 18), ("name", "效果中文名称", 18), ("description", "机制执行说明", 60), ("target", "默认作用目标", 16)]
    for col_idx, (k, label, width) in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=label)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = ALIGN_CENTER
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    for row_idx, eff in enumerate(EFFECT_DICT, 2):
        for col_idx, (k, _, _) in enumerate(headers, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=eff[k])
            cell.font = CELL_FONT
            cell.border = BORDER_THIN
            if k == "description":
                cell.alignment = ALIGN_LEFT
            else:
                cell.alignment = ALIGN_CENTER

# Create sheets
ws_base = wb.create_sheet(title="Cards_Base")
add_sheet_data(ws_base, "Cards_Base", BASE_CARDS)

ws_pvp = wb.create_sheet(title="Cards_PVP")
add_sheet_data(ws_pvp, "Cards_PVP", PVP_CARDS)

ws_pve = wb.create_sheet(title="Cards_PVE")
add_sheet_data(ws_pve, "Cards_PVE", PVE_CARDS)

ws_effects = wb.create_sheet(title="Card_Effects")
create_effects_sheet(ws_effects)

os.makedirs(os.path.dirname(OUTPUT_PATH), exist_ok=True)
wb.save(OUTPUT_PATH)
# Also copy to project root for user convenience
root_copy = os.path.join(PROJECT_ROOT, 'CardGame_Cards.xlsx')
wb.save(root_copy)

print(f"✅ Generated {OUTPUT_PATH} and {root_copy} successfully!")
print(f"   Cards_Base: {len(BASE_CARDS)} cards")
print(f"   Cards_PVP:  {len(PVP_CARDS)} cards")
print(f"   Cards_PVE:  {len(PVE_CARDS)} cards")
print(f"   Total:      {len(BASE_CARDS) + len(PVP_CARDS) + len(PVE_CARDS)} cards")
